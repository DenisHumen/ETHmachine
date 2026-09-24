"""Регрессии DeBank Checker: чужие портфели и мусор не попадают в баланс.

Что ловим:
  1. кошельку с реальными $0.08 записывались $859 000 — балансы собирались
     из ответов DeBank, не глядя, к какому адресу они относятся, и не
     сверялись с оценкой кошелька самой DeBank;
  2. аирдроп-спам (90% эмиссии токена «VITALIK» на одном кошельке)
     считался настоящим активом, потому что у DeBank высокий рейтинг
     бывает и у мусора;
  3. отчёт сортировал и суммировал по цифре вместе с мусором.

Браузер и сеть подменены заглушками, база и отчёт — во временном каталоге.
"""

from __future__ import annotations

import asyncio
import sqlite3

import pytest

from modules.debank import database as db
from modules.debank import debank_checker as dc

WALLET = "0xbB3bE20bcFa98f2d5880Ee1a255421111391Ebd1"
OTHER = "0xd8dA6BF26964aF9D7eEd9e03E53415D37aA96045"
API = "https://api.debank.com"


def _token(chain="eth", token_id="0x" + "ab" * 20, symbol="TKN", amount=1.0,
           price=1.0, score=1_000_000, supply=1e9, **flags):
    return {"chain": chain, "id": token_id, "symbol": symbol, "name": symbol,
            "amount": amount, "price": price, "credit_score": score,
            "total_supply": supply, **flags}


ETH = _token(token_id="eth", symbol="ETH", amount=0.00003, price=2700,
             score=44_000_000, supply=1.2e8)
WHALE_PORTFOLIO = [_token(symbol="USDC", amount=859_000, price=1.0, supply=5e10)]


class _Response:
    def __init__(self, path, owner, data, status=200):
        key = "id" if path in (dc.USER_PATH, dc.USED_CHAINS_PATH) else "user_addr"
        self.url = f"{API}{path}?{key}={owner.lower()}"
        self.status = status
        self._body = {"error_code": 0, "data": data}

    async def json(self):
        return self._body


def _profile(owner, usd_value):
    return _Response(dc.USER_PATH, owner,
                     {"user": {"desc": {"id": owner.lower(), "usd_value": usd_value}}})


def _chains(owner, chains):
    return _Response(dc.USED_CHAINS_PATH, owner, {"chains": chains})


def _cache(owner, tokens):
    return _Response(dc.CACHE_BALANCE_PATH, owner, tokens)


class _Page:
    """Страница-заглушка: на goto «загружает» заранее заданные ответы."""

    def __init__(self, responses):
        self._responses = responses
        self._handler = None

    def on(self, event, handler):
        assert event == "response"
        self._handler = handler

    async def goto(self, url, **kwargs):
        for response in self._responses:
            await self._handler(response)


class _Playwright:
    """Минимум от async_playwright, который нужен check_wallet_playwright."""

    def __init__(self, responses):
        page = _Page(responses)

        class _Browser:
            async def new_page(self):
                return page

            async def close(self):
                pass

        class _Chromium:
            async def launch(self, **kwargs):
                return _Browser()

        self.chromium = _Chromium()


@pytest.fixture()
def fast_retries(monkeypatch):
    """Одна попытка без пауз: ошибка видна сразу, тест не ждёт секундами."""
    monkeypatch.setattr(dc, "MAX_ATTEMPTS", 1)
    monkeypatch.setattr(dc, "SLEEP_BETWEEN_ACTIONS", (0, 0))


def _check(responses):
    async def run():
        return await dc.check_wallet_playwright(
            WALLET, None, asyncio.Semaphore(1), _Playwright(responses))
    return asyncio.run(run())


# ── Сбор данных ────────────────────────────────────────────────────────────

def test_response_owner_reads_both_query_keys():
    assert dc.response_owner(f"{API}/user?id={WALLET}") == WALLET.lower()
    assert dc.response_owner(
        f"{API}/token/balance_list?user_addr={WALLET}&chain=eth") == WALLET.lower()
    assert dc.response_owner(f"{API}/chain/list") == ""


def test_foreign_portfolio_is_ignored(fast_retries):
    """Ответ по чужому адресу не должен стать балансом проверяемого кошелька."""
    result = _check([
        _chains(WALLET, ["eth"]),
        _profile(WALLET, 0.08),
        _cache(WALLET, [ETH]),
        _cache(OTHER, WHALE_PORTFOLIO),     # чужой портфель перетёр бы свой
    ])

    assert result["success"], result["error"]
    assert result["foreign"] == 1
    assert [t["symbol"] for t in result["tokens"]] == ["ETH"]
    assert result["total_usd"] == pytest.approx(0.081)


def test_tokens_above_debank_net_worth_are_rejected(fast_retries):
    """Список токенов на $859k при оценке DeBank в $0.08 — не данные кошелька."""
    result = _check([
        _chains(WALLET, ["eth"]),
        _profile(WALLET, 0.08),
        _cache(WALLET, WHALE_PORTFOLIO),
    ])

    assert not result["success"]
    assert "DeBank оценивает кошелёк" in result["error"]
    assert result["tokens"] == []


def test_profile_of_another_address_is_rejected(fast_retries):
    wrong = _profile(WALLET, 0.08)
    wrong._body["data"]["user"]["desc"]["id"] = OTHER.lower()
    result = _check([_chains(WALLET, ["eth"]), wrong, _cache(WALLET, [ETH])])

    assert not result["success"]
    assert result["error"] == "DeBank отдал профиль другого адреса"


def test_empty_wallet_is_success(fast_retries):
    """Сетей нет — балансов DeBank не запрашивает, и это не ошибка."""
    result = _check([_chains(WALLET, []), _profile(WALLET, 0)])

    assert result["success"], result["error"]
    assert result["tokens"] == []


def test_defi_positions_do_not_trigger_mismatch():
    """Оценка DeBank больше токенов из-за DeFi — это нормально."""
    assert not dc.net_worth_mismatch(10.0, 5_000.0)
    assert not dc.net_worth_mismatch(1.0, 0.0)          # пыль в пределах допуска
    assert dc.net_worth_mismatch(859_000.0, 0.08)


# ── Мусор ──────────────────────────────────────────────────────────────────

def _parsed(**kwargs):
    return dc.parse_token_data([_token(**kwargs)])[0]


def test_airdrop_holding_big_share_of_supply_is_junk():
    """90% эмиссии на одном кошельке — спам, даже с высоким рейтингом."""
    token = _parsed(symbol="VITALIK", amount=900_000_000, price=0.00034,
                    supply=1e9, score=227_464)
    assert token["junk_reason"] and "эмиссии" in token["junk_reason"]


def test_low_credit_score_is_junk():
    token = _parsed(symbol="MEME", score=947)
    assert token["junk_reason"] and "рейтинг" in token["junk_reason"]


def test_real_token_is_not_junk():
    assert _parsed(symbol="ENS", amount=1144, price=7.1, supply=1e8,
                   score=434_867)["junk_reason"] is None


def test_native_coin_is_never_junk():
    """ETH в сети с дефолтным рейтингом и маленькой эмиссией — всё равно ETH."""
    token = _parsed(chain="zora", token_id="zora", symbol="ETH", amount=500,
                    price=2700, score=0, supply=10_000)
    assert token["junk_reason"] is None


def test_scam_flag_drops_token():
    assert dc.parse_token_data([_token(is_scam=True)]) == []


# ── База и отчёт ───────────────────────────────────────────────────────────

@pytest.fixture()
def tmp_db(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_DIR", tmp_path / "db")
    monkeypatch.setattr(db, "DB_FILE", tmp_path / "db" / "debank_checker.db")
    monkeypatch.setattr(dc, "project_root", tmp_path)
    db.init_database()
    return tmp_path


def _save(wallet, *tokens):
    db.save_token_balances_batch([
        (wallet, t["chain"], t["symbol"], t["name"], t["id"], t["amount"],
         t["price"], t["amount"] * t["price"], "", t["credit_score"],
         t["total_supply"])
        for t in tokens
    ])


def test_old_unverified_results_go_to_recheck(tmp_path, monkeypatch):
    """Результаты без сверки с DeBank могли быть чужими — их проверяем заново."""
    monkeypatch.setattr(db, "DB_DIR", tmp_path)
    monkeypatch.setattr(db, "DB_FILE", tmp_path / "debank_checker.db")
    with sqlite3.connect(tmp_path / "debank_checker.db") as conn:
        conn.execute("""CREATE TABLE debank_tasks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            wallet_address TEXT NOT NULL UNIQUE,
            status TEXT NOT NULL DEFAULT 'pending',
            attempts INTEGER DEFAULT 0, error_message TEXT,
            created_at TIMESTAMP, updated_at TIMESTAMP)""")
        conn.execute("INSERT INTO debank_tasks (wallet_address, status) "
                     "VALUES (?, 'completed')", (WALLET,))

    db.init_database()
    db.init_database()      # повторный вызов ничего не ломает

    assert [t["wallet_address"] for t in db.get_pending_tasks()] == [WALLET]
    task = db.get_all_tasks()[0]
    assert task["error_message"] == db.RECHECK_REASON


def test_report_counts_balance_without_junk(tmp_db):
    from openpyxl import load_workbook

    rich, poor, broken, fresh = (f"0x{n:040x}" for n in (1, 2, 3, 4))
    db.create_tasks([rich, poor, broken, fresh], {rich: "farm-1"})
    _save(rich,
          _token(token_id="eth", symbol="ETH", amount=1, price=2700),
          _token(symbol="VITALIK", amount=9e8, price=0.001, supply=1e9))
    _save(poor, _token(symbol="USDC", amount=5, price=1))
    _save(fresh, _token(symbol="USDC", amount=1_000, price=1))  # не проверен
    db.update_task_status(rich, "completed", net_worth_usd=902_700.0)
    db.update_task_status(poor, "completed", net_worth_usd=5.0)
    db.update_task_status(broken, "failed", "DeBank не отдал данные профиля")

    path = dc.save_results_xlsx([poor, rich, broken, fresh])
    wb = load_workbook(path)

    assert wb.sheetnames == ["Итоги", "Кошельки", "Токены", "Активы", "Сети",
                             "Мусор", "Не проверено"]

    wallets = [[c.value for c in row] for row in wb["Кошельки"].iter_rows(min_row=2)]
    # Сверху — самый богатый по балансу без мусора, имя из data.csv на месте.
    assert wallets[0][:5] == [2, "farm-1", rich, "готово", 2700]
    assert wallets[0][5] == pytest.approx(900_000)      # мусор отдельно
    assert wallets[1][2] == poor
    # Непроверенный кошелёк без цифр: его старым балансам не верим.
    assert [w[3] for w in wallets[2:]] == ["ошибка", "не проверен"]
    assert all(v is None for v in wallets[3][4:])

    tokens = {row[3].value for row in wb["Токены"].iter_rows(min_row=2)}
    assert tokens == {"ETH", "USDC"}
    junk = [row[3].value for row in wb["Мусор"].iter_rows(min_row=2)]
    assert junk == ["VITALIK"]

    summary = {row[0].value: row[1].value for row in wb["Итоги"].iter_rows(min_row=2)
               if row[0].value}
    assert summary["Баланс без мусора, $"] == pytest.approx(2705)
    assert summary["Не проверено"] == 1
