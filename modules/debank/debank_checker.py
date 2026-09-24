"""DeBank Balance Checker — балансы всех токенов кошелька во всех сетях.

Данные снимаются через браузер (Playwright): открывается профиль кошелька
на debank.com и перехватывается ответ ``cache_balance_list`` — прямой вызов
API упирается в anti-bot защиту.

Перехваченным ответам на слово не верим. Каждый должен относиться к
проверяемому адресу, а сумма токенов — сходиться с оценкой кошелька самой
DeBank (``/user``). Без этих проверок у кошельков с реальными $0.08 в отчёт
попадали $859 000 — чужой портфель вместо своего.

Настоящий баланс считается без мусорных аирдропов (``junk_reason``): DeBank
включает их в свою оценку, из-за чего суммы выглядели неправдоподобно.

Здесь же живут общие для обоих DeBank-модулей части: загрузка кошельков и
прокси, выбор источника кошельков, панель хода проверки и панель прогресса
по задачам — ``debank_protocol_checker`` импортирует их отсюда.
"""

import math
import asyncio
import time
import random
from collections import Counter
from pathlib import Path
from datetime import datetime
from urllib.parse import urlparse, parse_qs

from rich.console import Console
from rich.live import Live
from rich.panel import Panel
from rich.text import Text
from rich.table import Table
from rich.console import Group

from config.modules.general_config import (
    NUM_THREADS, RETRY_COUNT, SLEEP_BETWEEN_ACTIONS, DELAY_BETWEEN_ACCOUNTS
)
from config.modules.cfg_debank import (
    SKIP_SCAM_TOKENS, TRUSTED_CREDIT_SCORE, MAX_SUPPLY_SHARE, MIN_VALUE_USD,
    GRADIENT_MIN_USD, GRADIENT_MAX_USD
)
from modules.debank.database import (
    init_database, create_tasks, get_pending_tasks,
    update_task_status, save_token_balances_batch, delete_wallet_balances,
    get_all_balances, get_all_tasks, reset_database, get_task_statistics
)
from modules.simple_logger import logger
from modules.ui import ui
from modules.ui.menu_model import BACK_KEY, MenuItem, render_items

project_root = Path(__file__).parent.parent.parent

console = Console()

# Максимум параллельных браузерных контекстов
MAX_CONCURRENT = NUM_THREADS

# Пути ответов DeBank, из которых собираются балансы.
# ``cache_balance_list`` — агрегат сразу по всем сетям, ``balance_list`` — по
# одной сети (фронт дёргает его десятки раз). ``used_chains`` нужен, чтобы
# отличить пустой кошелёк от сорванной загрузки: если сетей у адреса нет,
# DeBank балансы вообще не запрашивает, и ждать их бессмысленно.
USED_CHAINS_PATH = '/user/used_chains'
CACHE_BALANCE_PATH = '/token/cache_balance_list'
CHAIN_BALANCE_PATH = '/token/balance_list'
# Профиль адреса: в ``desc.usd_value`` — оценка кошелька самой DeBank,
# та же, что в шапке профиля на сайте (токены + DeFi-позиции).
USER_PATH = '/user'

# Сколько секунд ждём данные после загрузки страницы.
DATA_WAIT_TIMEOUT = 30

# Каждая попытка — полный запуск браузера, поэтому ретраев немного:
# при RETRY_COUNT=15 один проблемный кошелёк занимал поток минут на десять.
MAX_ATTEMPTS = max(1, min(RETRY_COUNT, 3))

# Сверка собранных токенов с оценкой DeBank. В оценку входят ещё и
# DeFi-позиции, так что токенов может оказаться меньше, но не больше.
# Подмена чужим портфелем давала расхождение в тысячи раз, поэтому допуск
# щедрый: он гасит только разницу во времени между запросами.
NET_WORTH_TOLERANCE = 1.5
NET_WORTH_SLACK_USD = 5.0


# Источники кошельков: обе колонки лежат в общем data.csv проекта.
SOURCE_ADDRESSES = 'wallets'
SOURCE_PRIVATE_KEYS = 'private_keys'


def _legacy_lines(filename: str) -> list:
    """Строки из старых отдельных файлов (data/walletss.txt и т.п.).

    Нужны как запасной вариант, пока данные не перенесены в data.csv.
    """
    path = project_root / 'data' / filename
    try:
        with open(path, 'r', encoding='utf-8') as f:
            return [line.strip() for line in f
                    if line.strip() and not line.startswith('#')]
    except FileNotFoundError:
        return []
    except Exception as e:
        logger.error(f"Не удалось прочитать {path.name}: {e}")
        return []


def _address_from_key(private_key: str):
    """Адрес по приватному ключу. Сам ключ наружу не уходит."""
    from eth_account import Account

    key = private_key if private_key.startswith('0x') else '0x' + private_key
    try:
        return Account.from_key(key).address
    except Exception:
        return None


def load_wallet_rows(source: str = SOURCE_ADDRESSES) -> list:
    """Кошельки из общего data.csv — по адресам или по приватным ключам.

    Возвращает список ``{'address', 'proxy', 'name'}``: прокси берётся из
    той же строки, чтобы кошелёк всегда ходил через свой, а не через
    случайный; имя — чтобы в отчёте кошелёк можно было узнать.
    """
    from modules.data_manager import load_data

    rows = []
    invalid_keys = 0

    for row in load_data():
        proxy = (row.get('proxy') or '').strip()
        name = (row.get('name') or '').strip()
        if source == SOURCE_PRIVATE_KEYS:
            key = (row.get('private_key') or '').strip()
            if not key:
                continue
            address = _address_from_key(key)
            if not address:
                invalid_keys += 1      # сам ключ не логируем — это секрет
                continue
        else:
            address = (row.get('wallet_address') or '').strip()
            if not address.startswith('0x'):
                continue
        rows.append({'address': address, 'proxy': proxy, 'name': name})

    if invalid_keys:
        logger.warning(f"Пропущено невалидных ключей: {invalid_keys}")

    if rows:
        return rows

    # Нужной колонки в data.csv нет — поддерживаем старые отдельные файлы.
    if source == SOURCE_PRIVATE_KEYS:
        legacy = [a for a in
                  (_address_from_key(k) for k in _legacy_lines('private_keys.txt'))
                  if a]
        source_name = 'data/private_keys.txt'
    else:
        legacy = [a for a in _legacy_lines('walletss.txt') if a.startswith('0x')]
        source_name = 'data/walletss.txt'

    if legacy:
        logger.warning(f"В data.csv нужной колонки нет — взяли {source_name}")
    return [{'address': a, 'proxy': '', 'name': ''} for a in legacy]


def load_wallets() -> list:
    """Адреса из колонки wallet_address общего data.csv."""
    return [row['address'] for row in load_wallet_rows(SOURCE_ADDRESSES)]


def load_private_keys_as_wallets() -> list:
    """Адреса, выведенные из колонки private_key общего data.csv."""
    return [row['address'] for row in load_wallet_rows(SOURCE_PRIVATE_KEYS)]


def load_proxies() -> list:
    """Общий пул прокси: колонка proxy из data.csv, иначе старый proxy.csv."""
    from modules.data_manager import get_proxies

    proxies = [p for p in get_proxies() if '@' in p and ':' in p]
    if proxies:
        return proxies

    return [line for line in _legacy_lines('proxy.csv')
            if '@' in line and ':' in line
            and not line.lower().startswith(('proxy', 'login'))]


def parse_proxy_for_playwright(proxy_str: str) -> dict:
    """Конвертировать строку прокси в формат Playwright"""
    if not proxy_str:
        return None
    try:
        auth, addr = proxy_str.split('@')
        login, password = auth.split(':', 1)
        ip, port = addr.split(':', 1)
        return {
            'server': f'http://{ip}:{port}',
            'username': login,
            'password': password,
        }
    except Exception:
        return None


def is_scam_token(token: dict) -> bool:
    """Помечен ли токен скамом самой DeBank — такие выбрасываются сразу.

    Помечает она немногое: у спам-аирдропов флаги обычно такие же, как у
    ETH, их отсеивает уже ``junk_reason``.
    """
    return bool(SKIP_SCAM_TOKENS
                and (token.get('is_scam') or token.get('is_suspicious')))


def is_native(token: dict) -> bool:
    """Монета самой сети (ETH, BNB, POL…): у DeBank её id совпадает с id сети."""
    chain = str(token.get('chain') or '').lower()
    return bool(chain) and str(token.get('address') or '').lower() == chain


def junk_reason(token: dict):
    """Почему позиция не идёт в баланс; ``None`` — актив настоящий.

    ``token`` — позиция в формате ``parse_token_data``. Правила — в
    ``cfg_debank``; на данных vitalik.eth они оставляют от $1.16M по
    версии DeBank около $78k: ETH, ENS, USDT, UNI и прочие живые токены.
    """
    if is_native(token):
        return None

    supply = token.get('total_supply') or 0
    if supply > 0:
        share = (token.get('balance') or 0) / supply
        if share >= MAX_SUPPLY_SHARE:
            return f"{share:.1%} всей эмиссии на одном кошельке"

    score = token.get('credit_score') or 0
    if score < TRUSTED_CREDIT_SCORE:
        return f"рейтинг DeBank {score:,.0f} — ниже {TRUSTED_CREDIT_SCORE:,}"
    return None


def gradient_color(value_usd: float, min_usd: float = GRADIENT_MIN_USD,
                   max_usd: float = GRADIENT_MAX_USD):
    """Цвет заливки ячейки: чем дороже позиция, тем насыщеннее зелёный.

    Шкала логарифмическая: на фарм-кошельках почти всё стоит центы, и на
    линейной шкале разница между $1 и $100 была бы неразличима.
    ``None`` — сумма ниже порога, заливки нет.
    """
    if value_usd < min_usd:
        return None

    if value_usd >= max_usd:
        t = 1.0
    else:
        log_min = math.log10(max(min_usd, 0.01))
        log_max = math.log10(max(max_usd, 1.0))
        log_val = math.log10(max(value_usd, min_usd))
        span = log_max - log_min
        t = 0.0 if span <= 0 else max(0.0, min(1.0, (log_val - log_min) / span))

    # белый → светло-зелёный → насыщенный зелёный
    if t <= 0.5:
        s = t * 2
        r, g, b = (int(255 - 57 * s), int(255 - 16 * s), int(255 - 49 * s))
    else:
        s = (t - 0.5) * 2
        r, g, b = (int(198 - 159 * s), int(239 - 65 * s), int(206 - 110 * s))
    return f"{r:02X}{g:02X}{b:02X}"


def parse_token_data(data) -> list:
    """Парсинг данных токенов из API ответа DeBank"""
    tokens = []

    if isinstance(data, dict) and 'data' in data:
        token_list = data['data']
    elif isinstance(data, list):
        token_list = data
    else:
        return tokens

    if not isinstance(token_list, list):
        return tokens

    for token in token_list:
        if not isinstance(token, dict):
            continue
        amount = token.get('amount', 0)
        if not amount or float(amount) <= 0:
            continue
        if is_scam_token(token):
            continue

        price = token.get('price', 0) or 0
        parsed = {
            'chain': token.get('chain', ''),
            'symbol': token.get('symbol', token.get('optimized_symbol', 'UNKNOWN')),
            'name': token.get('name', ''),
            'address': token.get('id', token.get('contract_id', '')),
            'balance': float(amount),
            'price_usd': float(price),
            'value_usd': float(amount) * float(price),
            'logo_url': token.get('logo_url', ''),
            # Рейтинг токена у DeBank: у настоящих активов он огромный,
            # у свежих аирдропов — ноль.
            'credit_score': float(token.get('credit_score') or 0),
            # Эмиссия в той же сети — по ней видно аирдроп «10% выпуска».
            'total_supply': float(token.get('total_supply') or 0) or None,
        }
        parsed['junk_reason'] = junk_reason(parsed)
        tokens.append(parsed)
    return tokens


def chains_from_body(body):
    """Список сетей из ``/user/used_chains``; пустой список — кошелёк пустой."""
    if not isinstance(body, dict):
        return None
    data = body.get('data')
    if not isinstance(data, dict):
        return None
    chains = data.get('chains')
    return chains if isinstance(chains, list) else None


def response_owner(url: str) -> str:
    """Адрес, к которому относится запрос к API DeBank (в нижнем регистре).

    Адрес передаётся в query: ``user_addr`` у балансов и позиций, ``id`` у
    профиля и списка сетей. Пустая строка — запрос не про кошелёк.
    """
    query = parse_qs(urlparse(url).query)
    owner = (query.get('user_addr') or query.get('id') or [''])[0]
    return owner.strip().lower()


def profile_from_body(body):
    """``{'id', 'usd_value'}`` из ответа ``/user``; ``None`` — ответ не тот."""
    data = body.get('data') if isinstance(body, dict) else None
    user = data.get('user') if isinstance(data, dict) else None
    desc = user.get('desc') if isinstance(user, dict) else None
    if not isinstance(desc, dict) or not desc.get('id'):
        return None
    usd_value = desc.get('usd_value')
    return {
        'id': str(desc['id']).lower(),
        'usd_value': float(usd_value) if usd_value is not None else None,
    }


def watch_balance_api(page, wallet: str) -> dict:
    """Подписаться на ответы DeBank; словарь наполняется по ходу загрузки.

    Ответы по другим адресам отбрасываются и считаются в ``foreign``:
    собирать балансы, не глядя, чей это запрос, значит рано или поздно
    записать кошельку чужой портфель.
    """
    state = {'chains': None, 'cache': None, 'per_chain': {}, 'profile': None,
             'foreign': 0}
    wallet = wallet.lower()

    async def handle_response(response):
        url = response.url
        if 'api.debank.com' not in url:
            return
        parsed = urlparse(url)
        if parsed.path not in (USED_CHAINS_PATH, CACHE_BALANCE_PATH,
                               CHAIN_BALANCE_PATH, USER_PATH):
            return
        if response_owner(url) != wallet:
            state['foreign'] += 1
            return
        if response.status != 200:
            return
        try:
            if parsed.path == USER_PATH:
                profile = profile_from_body(await response.json())
                if profile is not None:
                    state['profile'] = profile
            elif parsed.path == USED_CHAINS_PATH:
                chains = chains_from_body(await response.json())
                if chains is not None:
                    state['chains'] = chains
            elif parsed.path == CACHE_BALANCE_PATH:
                data = (await response.json()).get('data')
                if isinstance(data, list):
                    state['cache'] = data
            elif parsed.path == CHAIN_BALANCE_PATH:
                data = (await response.json()).get('data')
                if isinstance(data, list):
                    chain = parse_qs(parsed.query).get('chain', [''])[0]
                    state['per_chain'][chain] = data
        except Exception:
            # Тело ответа могло стать недоступным — доберём из других запросов.
            pass

    page.on('response', handle_response)
    return state


def collected_tokens(state: dict) -> list:
    """Токены из агрегата, дополненные посетевыми ответами, без дублей."""
    merged = {}
    for token in (state.get('cache') or []):
        merged[(token.get('chain'), token.get('id'))] = token
    for tokens in state.get('per_chain', {}).values():
        for token in tokens:
            merged.setdefault((token.get('chain'), token.get('id')), token)
    return list(merged.values())


def balances_ready(state: dict) -> bool:
    """Пришло ли всё, что DeBank собирался отдать по этому кошельку."""
    chains = state.get('chains')
    if chains is None or state.get('profile') is None:
        return False        # профиль ещё не загрузился
    if not chains:
        return True         # сетей нет — кошелёк пустой, балансов не будет
    if state.get('cache') is not None:
        return True         # пришёл агрегат сразу по всем сетям
    return set(chains) <= set(state.get('per_chain', {}))


async def wait_for_balances(state: dict, timeout: float = DATA_WAIT_TIMEOUT) -> bool:
    """Ждём сами данные, а не фиксированную паузу: запросов на профиль десятки."""
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if balances_ready(state):
            return True
        await asyncio.sleep(0.5)
    return balances_ready(state)


def net_worth_mismatch(tokens_usd: float, net_worth) -> bool:
    """Токенов собрано больше, чем DeBank насчитала на весь кошелёк."""
    ceiling = max(net_worth or 0.0, 0.0) * NET_WORTH_TOLERANCE + NET_WORTH_SLACK_USD
    return tokens_usd > ceiling


def collection_error(state: dict, wallet: str, tokens: list, is_last: bool):
    """Что не так с собранными данными; ``None`` — им можно верить.

    Пустой кошелёк — не ошибка: сетей нет, значит и токенов нет.
    """
    if state['chains'] is None:
        # Нет даже списка сетей — блокировка, капча или прокси.
        return 'DeBank не отдал данные профиля'

    profile = state['profile']
    if profile is None:
        return 'DeBank не отдал оценку кошелька'
    if profile['id'] != wallet.lower():
        return 'DeBank отдал профиль другого адреса'

    # На последней попытке берём то, что успело прийти, — лучше неполный
    # список, чем никакого. Сверка ниже всё равно отсечёт чужие данные.
    if not balances_ready(state) and not (tokens and is_last):
        return 'Балансы не пришли за отведённое время'

    tokens_usd = sum(t['value_usd'] for t in tokens)
    if net_worth_mismatch(tokens_usd, profile['usd_value']):
        return (f"токенов на ${tokens_usd:,.2f}, а DeBank оценивает кошелёк "
                f"в ${profile['usd_value'] or 0:,.2f}")
    return None


def wallet_result(wallet: str, tokens: list, net_worth=None, error=None,
                  foreign: int = 0) -> dict:
    """Итог по кошельку: ``total_usd`` — баланс без мусора."""
    assets = [t for t in tokens if not t.get('junk_reason')]
    return {
        'wallet': wallet,
        'tokens': tokens,
        'assets': assets,
        'success': error is None,
        'total_usd': sum(t['value_usd'] for t in assets),
        'junk_usd': sum(t['value_usd'] for t in tokens if t.get('junk_reason')),
        'net_worth': net_worth,
        'foreign': foreign,
        'error': error,
    }


async def check_wallet_playwright(wallet: str, proxy_config: dict, semaphore: asyncio.Semaphore,
                                  playwright_instance) -> dict:
    """Проверка одного кошелька через Playwright browser"""
    async with semaphore:
        last_error = 'Max retries exceeded'
        foreign = 0

        for attempt in range(MAX_ATTEMPTS):
            try:
                launch_args = {'headless': True}
                if proxy_config:
                    launch_args['proxy'] = proxy_config

                browser = await playwright_instance.chromium.launch(**launch_args)
                try:
                    page = await browser.new_page()
                    state = watch_balance_api(page, wallet)

                    # Не networkidle: профиль тянет балансы десятками запросов
                    # и тишины в сети может не наступить вовсе.
                    await page.goto(
                        f'https://debank.com/profile/{wallet}',
                        wait_until='domcontentloaded',
                        timeout=45000
                    )

                    await wait_for_balances(state)

                    tokens = parse_token_data(collected_tokens(state))
                    foreign += state['foreign']
                    last_error = collection_error(
                        state, wallet, tokens, attempt == MAX_ATTEMPTS - 1)
                    if last_error is None:
                        return wallet_result(wallet, tokens,
                                             net_worth=state['profile']['usd_value'],
                                             foreign=foreign)

                finally:
                    await browser.close()

            except Exception as e:
                last_error = str(e)[:80]

            if attempt < MAX_ATTEMPTS - 1:
                await asyncio.sleep(random.uniform(*SLEEP_BETWEEN_ACTIONS))

    return wallet_result(wallet, [], error=last_error, foreign=foreign)


def progress_panel(icon: str, title: str, accent: str, total: int,
                   success: int, failed: int, logs: list) -> Panel:
    """Панель хода проверки для Live-вывода.

    Одна на оба DeBank-модуля: отличались только заголовок и цвет рамки.
    Живой прогресс рисует rich, а не UI-набор проекта: набор отдаёт
    статичные блоки, здесь же панель перерисовывается по ходу работы.
    """
    processed = success + failed
    percent = (processed / total * 100) if total > 0 else 0

    bar_width = 40
    filled = int(bar_width * percent / 100)
    bar_filled = "━" * filled
    bar_empty = "─" * (bar_width - filled)

    success_rate = (success / processed * 100) if processed > 0 else 100
    bar_style = "bright_green" if success_rate >= 80 else "yellow" if success_rate >= 50 else "red"

    table = Table(show_header=False, box=None, padding=(0, 1))
    table.add_column("L1", style="dim")
    table.add_column("V1")
    table.add_column("L2", style="dim")
    table.add_column("V2")
    table.add_row(
        "✅ Успешно:", Text(str(success), style="bold green"),
        "❌ Ошибки:", Text(str(failed), style="bold red")
    )
    table.add_row(
        "📊 Всего:", Text(str(total), style="bold"),
        "⚡ Параллельно:", Text(str(MAX_CONCURRENT), style="bold cyan")
    )

    header = Text()
    header.append(f"\n{icon} ", style="bold")
    header.append(f"{title}\n", style=f"bold {accent}")

    progress_bar = Text()
    progress_bar.append(bar_filled, style=f"bold {bar_style}")
    progress_bar.append(bar_empty, style="dim")
    progress_bar.append(f" {percent:.1f}%\n", style=f"bold {bar_style}")

    processed_text = Text()
    processed_text.append("Обработано: ", style="dim")
    processed_text.append(f"{processed}/{total}\n\n", style="bold")

    logs_section = Text()
    logs_section.append("📝 Последние события:\n", style="bold")
    for ts, msg, lvl in logs[-8:]:
        logs_section.append(f"  {ts} ", style="dim")
        style = {"SUCCESS": "green", "ERROR": "red", "WARNING": "yellow", "INFO": "cyan"}.get(lvl, "white")
        logs_section.append(f"{msg}\n", style=style)
    if not logs:
        logs_section.append("  Ожидание...\n", style="dim")

    content = Group(header, progress_bar, processed_text, table, Text(""), logs_section)

    return Panel(
        content,
        title=f"[bold {accent}]{icon} {title}[/bold {accent}]",
        subtitle="[dim]ETHmachine[/dim]",
        border_style=accent,
        padding=(1, 2)
    )


def balance_panel(total: int, success: int, failed: int, logs: list) -> Panel:
    return progress_panel("🏦", "DeBank · балансы токенов", "bright_blue",
                          total, success, failed, logs)


def choose_wallet_source() -> list | None:
    """Спрашивает, по какой колонке data.csv работать.

    Возвращает строки ``{'address', 'proxy'}`` — прокси нужен, чтобы
    кошелёк ходил через свой IP. ``None`` — пользователь вышел или
    подходящих кошельков не нашлось.
    """
    source = ui.menu("Откуда берём кошельки?", render_items([
        MenuItem(SOURCE_ADDRESSES, "Адреса кошельков",
                 "колонка wallet_address в data.csv", icon="📋"),
        MenuItem(SOURCE_PRIVATE_KEYS, "Приватные ключи",
                 "колонка private_key — в ней лежат секреты", icon="🔑"),
        MenuItem(BACK_KEY, "Назад", "", icon="←"),
    ]))
    if source in (None, BACK_KEY):
        return None

    rows = load_wallet_rows(source)
    if not rows:
        column = ('private_key' if source == SOURCE_PRIVATE_KEYS
                  else 'wallet_address')
        logger.error(f"В data.csv не заполнена колонка {column}")
        return None

    if source == SOURCE_PRIVATE_KEYS:
        logger.info(f"🔑 Из ключей получено адресов: {len(rows)}")
    else:
        logger.info(f"📋 Загружено кошельков: {len(rows)}")
    return rows


def wallet_proxy_map(rows: list) -> dict:
    """Карта «адрес → свой прокси» из строк data.csv."""
    return {row['address']: row['proxy'] for row in rows if row['proxy']}


def wallet_names(rows: list) -> dict:
    """Карта «адрес → имя из data.csv» — для отчёта."""
    return {row['address']: row['name'] for row in rows if row.get('name')}


def task_stats_panel(title: str, stats: dict) -> str | None:
    """Панель прогресса по задачам. ``None`` — в базе пока пусто."""
    labels = {"completed": "готово", "pending": "в очереди", "failed": "с ошибкой"}
    ordered = {label: stats.get(key, 0) for key, label in labels.items()}
    if not any(ordered.values()):
        return None
    ordered["total"] = sum(stats.values())
    return ui.stats_panel(title, ordered)


async def process_wallets_async(wallets: list, proxy_map: dict = None) -> dict:
    """Асинхронная обработка кошельков через Playwright.

    ``proxy_map`` — прокси из строки самого кошелька в data.csv; для
    кошельков без своего прокси берётся общий пул по кругу.
    """
    from playwright.async_api import async_playwright

    results = {}
    total = len(wallets)
    success = 0
    failed = 0
    logs = []

    proxies = load_proxies()
    proxy_map = proxy_map or {}
    paired = sum(1 for w in wallets if proxy_map.get(w))
    semaphore = asyncio.Semaphore(MAX_CONCURRENT)

    if not proxies:
        logger.warning("Прокси не найдены — запросы пойдут напрямую")

    delay_min, delay_max = DELAY_BETWEEN_ACCOUNTS
    logs.append((time.strftime("%H:%M:%S"), f"Запуск {total} кошельков ({MAX_CONCURRENT} параллельно, задержка {delay_min}-{delay_max}с)", "INFO"))
    if paired:
        logs.append((time.strftime("%H:%M:%S"),
                     f"У {paired} кошельков свой прокси из data.csv", "INFO"))
    if proxies:
        logs.append((time.strftime("%H:%M:%S"),
                     f"Общий пул прокси: {len(proxies)} (round-robin)", "INFO"))

    live = Live(balance_panel(total, 0, 0, logs), console=console, refresh_per_second=2)
    live.start()

    try:
        async with async_playwright() as p:

            async def process_wallet(idx, wallet):
                """Обработка одного кошелька с обновлением UI"""
                nonlocal success, failed
                # Свой прокси кошелька важнее общего пула: один адрес — один IP.
                proxy_str = proxy_map.get(wallet)
                if not proxy_str and proxies:
                    proxy_str = proxies[idx % len(proxies)]
                proxy_config = parse_proxy_for_playwright(proxy_str)

                short = f"{wallet[:6]}...{wallet[-4:]}"
                try:
                    result = await check_wallet_playwright(wallet, proxy_config, semaphore, p)
                    results[wallet] = result

                    if result['foreign']:
                        logs.append((
                            time.strftime("%H:%M:%S"),
                            f"[{short}] ⚠ отброшено ответов по чужим адресам: "
                            f"{result['foreign']}",
                            "WARNING"
                        ))

                    if result['success']:
                        success += 1
                        delete_wallet_balances(wallet)
                        if result['tokens']:
                            batch = [
                                (wallet, t['chain'], t['symbol'], t['name'],
                                 t['address'], t['balance'], t['price_usd'],
                                 t['value_usd'], t['logo_url'],
                                 t.get('credit_score', 0), t.get('total_supply'))
                                for t in result['tokens']
                            ]
                            save_token_balances_batch(batch)

                        update_task_status(wallet, 'completed',
                                           net_worth_usd=result['net_worth'])
                        count = len(result['assets'])
                        message = (f"[{short}] ✅ {count} "
                                   f"{plural(count, 'токен', 'токена', 'токенов')}"
                                   f" | ${result['total_usd']:,.2f}")
                        if result['junk_usd'] >= 0.01:
                            message += f" · мусор ${result['junk_usd']:,.2f}"
                        logs.append((time.strftime("%H:%M:%S"), message, "SUCCESS"))
                    else:
                        failed += 1
                        update_task_status(wallet, 'failed', result.get('error', 'Unknown'))
                        logs.append((
                            time.strftime("%H:%M:%S"),
                            f"[{short}] ❌ {result.get('error', 'Unknown')[:60]}",
                            "ERROR"
                        ))
                except Exception as e:
                    failed += 1
                    results[wallet] = wallet_result(wallet, [], error=str(e)[:100])
                    update_task_status(wallet, 'failed', str(e)[:100])
                    logs.append((time.strftime("%H:%M:%S"), f"[{short}] ❌ {str(e)[:30]}", "ERROR"))

                live.update(balance_panel(total, success, failed, logs))

            # Очередь и пул воркеров вместо создания всех задач заранее:
            # раньше между стартами ждали DELAY_BETWEEN_ACCOUNTS
            # последовательно, и на тысячах кошельков одна только раздача
            # задач растягивалась на часы, а реальная параллельность была
            # в разы ниже MAX_CONCURRENT.
            queue = asyncio.Queue()
            for item in enumerate(wallets):
                queue.put_nowait(item)

            async def worker():
                while True:
                    try:
                        idx, wallet = queue.get_nowait()
                    except asyncio.QueueEmpty:
                        return
                    await process_wallet(idx, wallet)
                    # Пауза между кошельками одного воркера.
                    await asyncio.sleep(random.uniform(*DELAY_BETWEEN_ACCOUNTS))

            workers = [asyncio.create_task(worker())
                       for _ in range(min(MAX_CONCURRENT, len(wallets)))]
            await asyncio.gather(*workers, return_exceptions=True)

    finally:
        live.stop()

    return results


def process_wallets(wallets: list, proxy_map: dict = None) -> dict:
    """Обёртка для запуска асинхронной обработки"""
    return asyncio.run(process_wallets_async(wallets, proxy_map))


# ── Отчёт ────────────────────────────────────────────────────────────────

STATUS_LABELS = {'completed': 'готово', 'failed': 'ошибка', 'pending': 'не проверен'}

# Сколько сетей перечислять в ячейке листа «Активы».
ASSET_CHAINS_SHOWN = 6


def stored_token(row: dict) -> dict:
    """Позиция из базы в формате ``parse_token_data``.

    Мусор размечается заново при каждом экспорте: пороги из cfg_debank
    применяются и к уже собранным данным, без повторной проверки.
    """
    token = {
        'chain': row['chain'],
        'symbol': row['token_symbol'],
        'name': row['token_name'] or '',
        'address': row['token_address'] or '',
        'balance': float(row['balance'] or 0),
        'price_usd': float(row['price_usd'] or 0),
        'value_usd': float(row['value_usd'] or 0),
        'credit_score': float(row.get('credit_score') or 0),
        'total_supply': row.get('total_supply'),
    }
    token['junk_reason'] = junk_reason(token)
    return token


def build_report(wallets: list, names: dict = None) -> list:
    """Строка на каждый кошелёк списка — из базы, а не из памяти запуска.

    В отчёт идут только кошельки выбранного списка: в базе могут лежать
    результаты по другому data-файлу. Балансы берутся только у проверенных
    кошельков — у остальных старые цифры не прошли сверку с DeBank.
    """
    names = names or {}
    tasks = {t['wallet_address']: t for t in get_all_tasks()}
    positions = {}
    for row in get_all_balances():
        positions.setdefault(row['wallet_address'], []).append(stored_token(row))

    report = []
    for index, wallet in enumerate(dict.fromkeys(wallets), start=1):
        task = tasks.get(wallet) or {}
        status = task.get('status') or 'pending'
        tokens = positions.get(wallet, []) if status == 'completed' else []
        assets = sorted((t for t in tokens if not t['junk_reason']),
                        key=lambda t: -t['value_usd'])
        junk = [t for t in tokens if t['junk_reason']]
        balance = sum(t['value_usd'] for t in assets)
        top = assets[0] if assets and assets[0]['value_usd'] > 0 else None
        report.append({
            'index': index,
            'wallet': wallet,
            'name': names.get(wallet) or task.get('account_name') or '',
            'status': status,
            'balance': balance,
            'junk': sum(t['value_usd'] for t in junk),
            'net_worth': task.get('net_worth_usd'),
            'assets': assets,
            'junk_tokens': junk,
            'chains': len({t['chain'] for t in assets}),
            'top': f"{top['symbol']} · {top['chain']}" if top else '',
            'top_share': top['value_usd'] / balance if top and balance > 0 else None,
            'error': task.get('error_message') or '',
            'attempts': task.get('attempts') or 0,
            'updated_at': str(task.get('updated_at') or '')[:19].replace('T', ' '),
        })
    return report


def save_results_xlsx(wallets: list, names: dict = None):
    """Excel-отчёт по списку кошельков.

    Листы: «Итоги», «Кошельки», «Токены», «Активы», «Сети», «Мусор» и
    «Не проверено». Каждый отвечает на свой вопрос: сколько всего, у кого,
    из чего складывается, чего и сколько в сумме, где лежит, что отброшено
    и кого ещё надо добрать. Главная цифра везде — баланс без мусора;
    оценка DeBank стоит рядом, чтобы сверить с сайтом.
    """
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    rows = build_report(wallets, names)
    if not rows:
        logger.warning("Список кошельков пуст — экспортировать нечего")
        return None

    result_dir = project_root / 'result'
    result_dir.mkdir(exist_ok=True)

    # ── Стили ──
    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="2C3E50", end_color="2C3E50")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    edge = Side(style='thin', color='D5D8DC')
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    mono_font = Font(name='Consolas', size=10)
    body_font = Font(size=10)
    money_font = Font(bold=True, size=10)
    muted_font = Font(size=10, color="7F8C8D")
    junk_fill = "FDEBD0"        # мягкий оранжевый — «на эту сумму не рассчитывай»
    error_fill = "FADBD8"

    MONEY = '$#,##0.00'
    PRICE = '$#,##0.00######'
    AMOUNT = '#,##0.########'
    SHARE = '0.0%'

    def put_header(ws, columns):
        """``columns`` — ``(заголовок, ширина, пояснение)``; пояснение
        всплывает при наведении на заголовок."""
        for idx, (title, width, note) in enumerate(columns, 1):
            cell = ws.cell(row=1, column=idx, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border
            if note:
                cell.comment = Comment(note, "ETHmachine", width=260, height=110)
            ws.column_dimensions[cell.column_letter].width = width
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = 'A2'

    def put(ws, row, col, value, *, font=None, fmt=None, fill=None,
            align=None, wrap=False):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = font or body_font
        cell.border = border
        if fmt:
            cell.number_format = fmt
        if fill:
            cell.fill = PatternFill("solid", start_color=fill, end_color=fill)
        if align or wrap:
            cell.alignment = Alignment(horizontal=align, vertical="top",
                                       wrap_text=wrap)
        return cell

    def put_money(ws, row, col, value):
        put(ws, row, col, round(value, 2), font=money_font, fmt=MONEY,
            fill=gradient_color(value))

    def finish(ws):
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions

    checked = [r for r in rows if r['status'] == 'completed']
    failed = [r for r in rows if r['status'] == 'failed']
    pending = [r for r in rows if r['status'] == 'pending']
    total_balance = sum(r['balance'] for r in checked)
    total_junk = sum(r['junk'] for r in checked)
    total_debank = sum(r['net_worth'] or 0 for r in checked)

    # Порядок для листов с кошельками: проверенные по балансу, затем с
    # ошибкой, затем не проверенные; внутри — как в data.csv.
    status_order = {'completed': 0, 'failed': 1, 'pending': 2}
    ordered = sorted(rows, key=lambda r: (status_order.get(r['status'], 3),
                                          -r['balance'], r['index']))

    wb = Workbook()

    # ── Итоги ──
    ws = wb.active
    ws.title = "Итоги"
    put_header(ws, [("Показатель", 38, None), ("Значение", 22, None)])
    ws.freeze_panes = None
    summary = [
        ("Дата отчёта", datetime.now().strftime("%d.%m.%Y %H:%M"), None),
        ("Кошельков в списке", len(rows), None),
        ("Проверено", len(checked), None),
        ("  из них с балансом от $1", sum(1 for r in checked if r['balance'] >= 1), None),
        ("С ошибкой", len(failed), None),
        ("Не проверено", len(pending), None),
        ("Баланс без мусора, $", total_balance, MONEY),
        ("Мусор (не учтён), $", total_junk, MONEY),
        ("Оценка DeBank, $", total_debank, MONEY),
    ]
    for line, (label, value, fmt) in enumerate(summary, start=2):
        put(ws, line, 1, label)
        if fmt:
            put(ws, line, 2, round(value, 2), font=money_font, fmt=fmt)
        else:
            put(ws, line, 2, value, align="right")
    notes = [
        "Баланс — токены без мусора: то, что реально есть на кошельках.",
        "Оценка DeBank — как на сайте: с DeFi-позициями и мусором.",
        f"Мусор — рейтинг токена у DeBank ниже {TRUSTED_CREDIT_SCORE:,} или "
        f"больше {MAX_SUPPLY_SHARE:.0%} всей эмиссии на одном кошельке. "
        "Монеты сетей (ETH, BNB, POL…) мусором не бывают. "
        "Пороги — в config/modules/cfg_debank.py.",
    ]
    for offset, note in enumerate(notes, start=len(summary) + 3):
        ws.merge_cells(start_row=offset, start_column=1, end_row=offset, end_column=2)
        cell = ws.cell(row=offset, column=1, value=note)
        cell.font = muted_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[offset].height = 30 if len(note) < 90 else 58

    # ── Кошельки ──
    ws = wb.create_sheet("Кошельки")
    put_header(ws, [
        ("№", 7, "Номер кошелька в списке. Сортировка по нему вернёт порядок data.csv"),
        ("Имя", 16, None),
        ("Кошелёк", 46, None),
        ("Статус", 12, None),
        ("Баланс, $", 15, "Токены без мусора — то, что реально есть на кошельке"),
        ("Мусор, $", 13, "Аирдропы без ликвидности: DeBank их оценивает, но "
                         "продать их нельзя. Причины — на листе «Мусор»"),
        ("DeBank, $", 14, "Оценка кошелька на сайте DeBank: с DeFi-позициями и мусором"),
        ("Токенов", 9, "Сколько настоящих токенов, без мусора"),
        ("Сетей", 8, None),
        ("Крупнейший актив", 24, None),
        ("Его доля", 10, None),
    ])
    for line, row in enumerate(ordered, start=2):
        is_done = row['status'] == 'completed'
        put(ws, line, 1, row['index'], align="center")
        put(ws, line, 2, row['name'])
        put(ws, line, 3, row['wallet'], font=mono_font)
        put(ws, line, 4, STATUS_LABELS.get(row['status'], row['status']),
            align="center", font=None if is_done else muted_font,
            fill=error_fill if row['status'] == 'failed' else None)
        if not is_done:
            for col in range(5, 12):
                put(ws, line, col, None)
            continue
        put_money(ws, line, 5, row['balance'])
        put(ws, line, 6, round(row['junk'], 2), fmt=MONEY,
            fill=junk_fill if row['junk'] >= GRADIENT_MIN_USD else None)
        net_worth = row['net_worth']
        put(ws, line, 7, round(net_worth, 2) if net_worth is not None else None,
            fmt=MONEY)
        put(ws, line, 8, len(row['assets']), align="center")
        put(ws, line, 9, row['chains'], align="center")
        put(ws, line, 10, row['top'])
        put(ws, line, 11, row['top_share'], fmt=SHARE, align="center")
    finish(ws)

    # ── Токены: все настоящие позиции ──
    ws = wb.create_sheet("Токены")
    put_header(ws, [
        ("Кошелёк", 46, None), ("Имя", 16, None), ("Сеть", 12, None),
        ("Токен", 14, None), ("Название", 26, None), ("Количество", 22, None),
        ("Цена, $", 15, None), ("Стоимость, $", 15, None),
    ])
    line = 1
    skipped_dust = 0
    for row in ordered:
        for t in row['assets']:
            if t['value_usd'] < MIN_VALUE_USD:
                skipped_dust += 1
                continue
            line += 1
            put(ws, line, 1, row['wallet'], font=mono_font)
            put(ws, line, 2, row['name'])
            put(ws, line, 3, t['chain'], align="center")
            put(ws, line, 4, t['symbol'])
            put(ws, line, 5, t['name'])
            put(ws, line, 6, t['balance'], fmt=AMOUNT)
            put(ws, line, 7, t['price_usd'], fmt=PRICE)
            put_money(ws, line, 8, t['value_usd'])
    finish(ws)

    # ── Активы: сколько чего всего, по всем кошелькам и сетям ──
    assets = {}
    for row in checked:
        for t in row['assets']:
            stat = assets.setdefault(t['symbol'], {
                'chains': {}, 'wallets': set(), 'amount': 0.0, 'value': 0.0})
            stat['chains'][t['chain']] = stat['chains'].get(t['chain'], 0.0) + t['value_usd']
            stat['wallets'].add(row['wallet'])
            stat['amount'] += t['balance']
            stat['value'] += t['value_usd']

    ws = wb.create_sheet("Активы")
    put_header(ws, [
        ("Токен", 14, None),
        ("Кошельков", 11, None),
        ("Количество", 22, "Сумма по всем кошелькам и сетям"),
        ("Стоимость, $", 16, None),
        ("Доля", 9, "Доля в общем балансе без мусора"),
        ("Сетей", 8, None),
        ("Где лежит", 44, "Сети по убыванию стоимости"),
    ])
    for line, (symbol, stat) in enumerate(
            sorted(assets.items(), key=lambda kv: -kv[1]['value']), start=2):
        chains = [c for c, _ in sorted(stat['chains'].items(), key=lambda kv: -kv[1])]
        where = ", ".join(chains[:ASSET_CHAINS_SHOWN])
        if len(chains) > ASSET_CHAINS_SHOWN:
            where += f" и ещё {len(chains) - ASSET_CHAINS_SHOWN}"
        put(ws, line, 1, symbol)
        put(ws, line, 2, len(stat['wallets']), align="center")
        put(ws, line, 3, stat['amount'], fmt=AMOUNT)
        put_money(ws, line, 4, stat['value'])
        put(ws, line, 5, stat['value'] / total_balance if total_balance else 0,
            fmt=SHARE, align="center")
        put(ws, line, 6, len(chains), align="center")
        put(ws, line, 7, where)
    finish(ws)

    # ── Сети: где лежат деньги ──
    chains = {}
    for row in checked:
        for t in row['assets'] + row['junk_tokens']:
            stat = chains.setdefault(t['chain'], {
                'wallets': set(), 'tokens': 0, 'value': 0.0, 'junk': 0.0})
            if t['junk_reason']:
                stat['junk'] += t['value_usd']
            else:
                stat['wallets'].add(row['wallet'])
                stat['tokens'] += 1
                stat['value'] += t['value_usd']

    ws = wb.create_sheet("Сети")
    put_header(ws, [
        ("Сеть", 14, None),
        ("Кошельков", 11, "Кошельки с настоящими токенами в этой сети"),
        ("Токенов", 10, None),
        ("Баланс, $", 16, None),
        ("Доля", 9, "Доля в общем балансе без мусора"),
        ("Мусор, $", 13, None),
    ])
    for line, (chain, stat) in enumerate(
            sorted(chains.items(), key=lambda kv: (-kv[1]['value'], -kv[1]['junk'])),
            start=2):
        put(ws, line, 1, chain)
        put(ws, line, 2, len(stat['wallets']), align="center")
        put(ws, line, 3, stat['tokens'], align="center")
        put_money(ws, line, 4, stat['value'])
        put(ws, line, 5, stat['value'] / total_balance if total_balance else 0,
            fmt=SHARE, align="center")
        put(ws, line, 6, round(stat['junk'], 2), fmt=MONEY,
            fill=junk_fill if stat['junk'] >= GRADIENT_MIN_USD else None)
    finish(ws)

    # ── Мусор: что отброшено и почему ──
    junk_rows = sorted(
        ((row, t) for row in checked for t in row['junk_tokens']
         if t['value_usd'] >= 0.01),
        key=lambda pair: -pair[1]['value_usd'])
    ws = wb.create_sheet("Мусор")
    put_header(ws, [
        ("Кошелёк", 46, None), ("Имя", 16, None), ("Сеть", 12, None),
        ("Токен", 14, None), ("Количество", 22, None), ("Цена, $", 15, None),
        ("Оценка DeBank, $", 16, "Во столько DeBank оценивает позицию. "
                                 "В баланс она не входит"),
        ("Почему мусор", 44, None),
    ])
    for line, (row, t) in enumerate(junk_rows, start=2):
        put(ws, line, 1, row['wallet'], font=mono_font)
        put(ws, line, 2, row['name'])
        put(ws, line, 3, t['chain'], align="center")
        put(ws, line, 4, t['symbol'])
        put(ws, line, 5, t['balance'], fmt=AMOUNT)
        put(ws, line, 6, t['price_usd'], fmt=PRICE)
        put(ws, line, 7, round(t['value_usd'], 2), fmt=MONEY, fill=junk_fill)
        put(ws, line, 8, t['junk_reason'])
    finish(ws)

    # ── Не проверено: кого ещё добрать ──
    if failed or pending:
        ws = wb.create_sheet("Не проверено")
        put_header(ws, [
            ("№", 7, None), ("Имя", 16, None), ("Кошелёк", 46, None),
            ("Статус", 12, None), ("Попыток", 9, None),
            ("Причина", 60, "«Продолжить проверку» в меню DeBank доберёт эти кошельки"),
            ("Обновлено", 19, None),
        ])
        for line, row in enumerate(failed + pending, start=2):
            put(ws, line, 1, row['index'], align="center")
            put(ws, line, 2, row['name'])
            put(ws, line, 3, row['wallet'], font=mono_font)
            put(ws, line, 4, STATUS_LABELS.get(row['status'], row['status']),
                align="center",
                fill=error_fill if row['status'] == 'failed' else None)
            put(ws, line, 5, row['attempts'], align="center")
            put(ws, line, 6, row['error'], wrap=True)
            put(ws, line, 7, row['updated_at'], align="center")
        finish(ws)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filepath = result_dir / f"debank_balances_{timestamp}.xlsx"
    wb.save(str(filepath))

    logger.success(f"Отчёт сохранён: {filepath}")
    logger.info(f"Кошельков: {len(rows)} · проверено: {len(checked)} · "
                f"с ошибкой: {len(failed)} · не проверено: {len(pending)}")
    logger.info(f"Баланс без мусора: ${total_balance:,.2f} · "
                f"мусор (не учтён): ${total_junk:,.2f} · "
                f"оценка DeBank: ${total_debank:,.2f}")
    if skipped_dust:
        logger.info(f"Скрыто позиций дешевле ${MIN_VALUE_USD:.2f}: {skipped_dust}")
    if failed or pending:
        logger.warning("Не все кошельки проверены — список на листе «Не проверено», "
                       "«Продолжить проверку» доберёт их")
    return filepath


def plural(count: int, one: str, few: str, many: str) -> str:
    """Русское склонение: 1 токен, 2 токена, 5 токенов."""
    tail = abs(count) % 100
    if 11 <= tail <= 14:
        return many
    tail %= 10
    if tail == 1:
        return one
    if 2 <= tail <= 4:
        return few
    return many


def top_wallets_panel(title: str, results: dict, count_key: str,
                      forms: tuple) -> str | None:
    """Топ-10 кошельков по стоимости. ``None`` — ни на одном ничего нет."""
    top = sorted(
        (r for r in results.values()
         if r['success'] and r.get('total_usd', 0) > 0),
        key=lambda r: r['total_usd'], reverse=True,
    )[:10]
    if not top:
        return None

    amounts = [f"${r['total_usd']:,.2f}" for r in top]
    width = max(len(amount) for amount in amounts)
    lines = []
    for index, (row, amount) in enumerate(zip(top, amounts), 1):
        count = len(row[count_key])
        lines.append(
            f"{index:>2}. {ui.pad(ui.shorten_address(row['wallet'], 10, 6), 18)}"
            f"{ui.theme.FG_OK}{ui.pad(amount, width, 'right')}{ui.theme.RESET}  "
            f"{ui.theme.FG_MUTED}{count} {plural(count, *forms)}{ui.theme.RESET}"
        )
    return ui.panel(title, lines)


def print_summary(results: dict):
    """Итоги проверки балансов: главная цифра — баланс без мусора."""
    succeeded = [r for r in results.values() if r['success']]
    failed = [r for r in results.values() if not r['success']]

    ui.print_lines(ui.stats_panel("Итоги проверки балансов", {
        "проверено": len(succeeded),
        "с ошибкой": len(failed),
        "токенов без мусора": sum(len(r['assets']) for r in succeeded),
        "баланс без мусора": f"${sum(r['total_usd'] for r in succeeded):,.2f}",
        "мусор, не учтён": f"${sum(r['junk_usd'] for r in succeeded):,.2f}",
        "total": len(results),
    }))

    top = top_wallets_panel("Топ кошельков по балансу без мусора", results,
                            "assets", ("токен", "токена", "токенов"))
    if top:
        ui.print_lines(top)

    if failed:
        reasons = Counter(r['error'] for r in failed).most_common(3)
        ui.print_lines(ui.panel("Частые ошибки", [
            f"{count:>5} × {error}" for error, count in reasons]))

    foreign = sum(r.get('foreign', 0) for r in results.values())
    if foreign:
        logger.warning(f"Отброшено ответов DeBank по чужим адресам: {foreign} — "
                       f"в балансы они не попали")


def debank_checker_menu():
    """Меню проверки балансов — точка входа из главного меню."""
    rows = choose_wallet_source()
    if not rows:
        return

    wallets = [row["address"] for row in rows]
    proxy_map = wallet_proxy_map(rows)
    names = wallet_names(rows)

    init_database()

    stats = task_stats_panel("Прогресс · балансы", get_task_statistics())
    if stats:
        ui.print_lines(stats)

    action = ui.menu("Проверка балансов DeBank", render_items([
        MenuItem("continue", "Продолжить проверку",
                 "взять из базы незавершённые кошельки", icon="▶️"),
        MenuItem("reset", "Начать заново",
                 "очистить базу и проверить все кошельки", icon="🔄"),
        MenuItem("export", "Экспорт в Excel",
                 "итоги, кошельки, токены, активы, сети, мусор", icon="📊"),
        MenuItem(BACK_KEY, "Назад", "", icon="←"),
    ]))
    if action in (None, BACK_KEY):
        return

    if action == "export":
        save_results_xlsx(wallets, names)
        return

    if action == "reset":
        reset_database()
        logger.warning("База балансов очищена")
    create_tasks(wallets, names)

    # Только кошельки выбранного списка: в базе могут висеть задачи
    # по другому data-файлу.
    in_list = set(wallets)
    wallets_to_check = [t['wallet_address'] for t in get_pending_tasks()
                        if t['wallet_address'] in in_list]
    if not wallets_to_check:
        logger.info("Все кошельки уже проверены — для повторной проверки "
                    "выберите «Начать заново»")
        save_results_xlsx(wallets, names)
        return

    logger.info(f"📋 Задач к выполнению: {len(wallets_to_check)}")

    results = process_wallets(wallets_to_check, proxy_map)
    print_summary(results)
    save_results_xlsx(wallets, names)
    logger.success("Проверка балансов DeBank завершена")


__all__ = [
    "debank_checker_menu",
    # Общее с debank_protocol_checker.
    "load_wallets", "load_private_keys_as_wallets", "load_proxies",
    "parse_proxy_for_playwright", "choose_wallet_source", "wallet_proxy_map",
    "wallet_names", "load_wallet_rows", "progress_panel", "response_owner",
    "task_stats_panel", "top_wallets_panel", "plural",
]
